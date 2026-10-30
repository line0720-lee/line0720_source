"""결과 엑셀 생성·서식·건별 저장.

실행 1회당 xlsx 1개를 만들고 파일 1건마다 저장한다. 파일이 열려 있어 저장하지 못하면 3회 재시도한 뒤
'_임시저장.xlsx'로 바꿔 저장하고 알림을 보낸다. 셀 한도(32,767자)를 넘는 변환 내용은 '전문' 시트와 txt로 보존한다.
저장은 임시 파일(.xlsx.tmp)에 쓴 뒤 바꿔치기하므로 저장 도중 프로그램이 끝나도 이전 행은 잃지 않는다.
"""
from __future__ import annotations

import logging
import os
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from records import ResultRow

logger = logging.getLogger(__name__)

CELL_LIMIT = 32767  # 엑셀 셀 1칸에 넣을 수 있는 최대 글자 수
OVERFLOW_NOTE = "\n…(이하 생략: '전문' 시트와 txt 파일 참조)"
OVERFLOW_NOTE_NO_TXT = "\n…(이하 생략: txt 파일은 저장하지 못했습니다. 전체 내용은 '전문' 시트에만 있습니다)"

MAIN_SHEET = "STT결과"
FULLTEXT_SHEET = "전문"
_MAIN_HEADERS = ["번호", "파일명", "변환 내용", "재생 시간", "처리 시간(초)", "상태", "변환 일시"]
_MAIN_WIDTHS = {"A": 6, "B": 40, "C": 100, "D": 12, "E": 14, "F": 28, "G": 20}
_FULLTEXT_HEADERS = ["번호", "파일명", "조각", "내용"]
_FULLTEXT_WIDTHS = {"A": 6, "B": 40, "C": 8, "D": 100}

_SAVE_ATTEMPTS = 3
_RETRY_INTERVAL_SEC = 1
_FALLBACK_SUFFIX = "_임시저장"

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")
_FAILED_FONT = Font(color="FF0000")
_TOP = Alignment(vertical="top")
_TOP_WRAP = Alignment(vertical="top", wrap_text=True)


def format_duration(seconds: float) -> str:
    """초를 hh:mm:ss로(반올림). 24시간 이상은 시간을 그대로 늘려 쓴다(예: 25:00:00)."""
    total = int(seconds + 0.5)
    hours, rest = divmod(total, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _clean(text: str) -> str:
    """엑셀이 허용하지 않는 제어 문자를 제거한다(저장 오류 방지)."""
    return ILLEGAL_CHARACTERS_RE.sub("", text)


def _set_text(cell, value: str, alignment: Alignment) -> None:
    """문자열 셀 기록. '='로 시작해도 수식이 아니라 글자로 저장한다."""
    cell.value = value
    if value.startswith("="):
        cell.data_type = "s"
    cell.alignment = alignment


def _write_header(ws, headers: list[str], widths: dict[str, int]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


class ExcelWriter:
    """batch_runner.RowWriter를 구현한다. 생성 즉시 머리글만 있는 파일을 저장하고, add_row마다 저장한다."""

    def __init__(
        self,
        result_dir: Path,
        started_at: datetime,
        on_warning: Callable[[str], None] | None = None,
        *,
        sleep: Callable[[float], None] = time.sleep,
    ) -> None:
        self._result_dir = Path(result_dir)
        self._started_at = started_at
        self._on_warning = on_warning
        self._sleep = sleep
        self._fell_back = False
        self._next_row = 2
        self._fulltext_ws = None
        self._fulltext_next_row = 2
        self.path = self._result_dir / f"STT_결과_{started_at:%Y%m%d_%H%M%S}.xlsx"

        self._wb = Workbook()
        self._ws = self._wb.active
        self._ws.title = MAIN_SHEET
        _write_header(self._ws, _MAIN_HEADERS, _MAIN_WIDTHS)
        self._save()

    def add_row(self, row: ResultRow) -> None:
        """결과 1건을 기록하고 바로 저장한다. 저장 실패 처리는 _save 참조."""
        clean = _clean(row.text)  # 셀·'전문' 시트·txt 모두 같은 제거본을 쓴다(한 번만 계산)
        if len(clean) > CELL_LIMIT:
            note = OVERFLOW_NOTE if self._store_full_text(row, clean) else OVERFLOW_NOTE_NO_TXT
            shown = clean[: CELL_LIMIT - len(note)] + note
        else:
            shown = clean

        r = self._next_row
        self._next_row += 1
        ws = self._ws
        ws.cell(row=r, column=1, value=row.index).alignment = _TOP
        _set_text(ws.cell(row=r, column=2), _clean(row.name), _TOP)
        _set_text(ws.cell(row=r, column=3), shown, _TOP_WRAP)
        duration_cell = ws.cell(row=r, column=4)
        duration_cell.value = None if row.duration_sec is None else format_duration(row.duration_sec)
        duration_cell.alignment = _TOP
        ws.cell(row=r, column=5, value=round(row.elapsed_sec)).alignment = _TOP
        status_cell = ws.cell(row=r, column=6)
        _set_text(status_cell, _clean(row.status), _TOP)
        if row.failed:
            status_cell.font = _FAILED_FONT
        ws.cell(row=r, column=7, value=row.finished_at.strftime("%Y-%m-%d %H:%M:%S")).alignment = _TOP
        self._save()

    # --- 셀 한도 초과 처리 -------------------------------------------------

    def _store_full_text(self, row: ResultRow, full: str) -> bool:
        """'전문' 시트에 한도 크기로 나누어 기록하고 같은 이름의 txt도 저장한다.

        full은 제어 문자를 제거한 내용이다(txt도 같은 내용 → txt == 시트 조각을 이어 붙인 것).
        txt를 쓰지 못하면(경로가 너무 김·디스크 부족 등) 경고만 남기고 False를 돌려준다.
        전체 내용은 '전문' 시트에 이미 있으므로 행 기록과 일괄 처리는 계속한다.
        """
        if self._fulltext_ws is None:
            self._fulltext_ws = self._wb.create_sheet(FULLTEXT_SHEET)
            _write_header(self._fulltext_ws, _FULLTEXT_HEADERS, _FULLTEXT_WIDTHS)
        ws = self._fulltext_ws
        for piece, start in enumerate(range(0, len(full), CELL_LIMIT), start=1):
            r = self._fulltext_next_row
            self._fulltext_next_row += 1
            ws.cell(row=r, column=1, value=row.index).alignment = _TOP
            _set_text(ws.cell(row=r, column=2), _clean(row.name), _TOP)
            ws.cell(row=r, column=3, value=piece).alignment = _TOP
            _set_text(ws.cell(row=r, column=4), full[start : start + CELL_LIMIT], _TOP_WRAP)

        txt_name = row.name
        try:
            txt_path = self._free_txt_path(row.name)
            txt_name = txt_path.name
            # utf-8-sig: 윈도우 메모장에서도 한글이 깨지지 않도록 BOM을 붙인다
            txt_path.write_text(full, encoding="utf-8-sig")
        except OSError:
            message = f"전문 txt 파일을 저장하지 못했습니다: {txt_name} ('전문' 시트에는 저장됨)"
            logger.exception("%s (원본: %s)", message, row.name)
            if self._on_warning is not None:
                self._on_warning(message)
            return False
        logger.info("셀 한도 초과: '%s' 전문을 %s에 저장", row.name, txt_path.name)
        return True

    def _free_txt_path(self, name: str) -> Path:
        """B열 이름에서 txt 경로를 만든다. result 폴더를 재사용하므로 이미 있으면 덮어쓰지 않고 실행 시각을 붙인다."""
        flat = name.replace("\\", "_")
        stem = re.sub(r"\.wav$", "", flat, flags=re.IGNORECASE)
        candidate = self._result_dir / f"{stem}.txt"
        if not candidate.exists():
            return candidate
        stamped = f"{stem}_{self._started_at:%Y%m%d_%H%M%S}"
        candidate = self._result_dir / f"{stamped}.txt"
        number = 2
        while candidate.exists():  # 한 실행 안에서 이름이 겹치는 경우까지 덮어쓰지 않는다
            candidate = self._result_dir / f"{stamped}_{number}.txt"
            number += 1
        return candidate

    # --- 저장 ---------------------------------------------------------------

    def _save(self) -> None:
        """현재 경로에 저장. 파일이 열려 있으면(PermissionError) 3회 시도 후 '_임시저장.xlsx'로 옮겨 저장한다."""
        try:
            self._save_with_retry(self.path)
            return
        except PermissionError:
            if self._fell_back:
                raise  # 대체 파일마저 저장하지 못하면 더 할 수 있는 일이 없다
        fallback = self.path.with_name(f"{self.path.stem}{_FALLBACK_SUFFIX}{self.path.suffix}")
        message = f"결과 파일이 열려 있어 저장하지 못했습니다. 다른 이름으로 저장합니다: {fallback.name}"
        logger.warning("%s (원본: %s)", message, self.path)
        self._write_file(fallback)  # 실패하면 예외가 그대로 전달되어 일괄 처리가 멈춘다
        self.path = fallback
        self._fell_back = True
        if self._on_warning is not None:
            self._on_warning(message)

    def _write_file(self, target: Path) -> None:
        """원자적 저장: 같은 폴더의 '<이름>.xlsx.tmp'에 다 쓴 뒤 os.replace로 한 번에 바꾼다.

        저장 도중 강제 종료(os._exit)·전원 꺼짐·디스크 부족이 나도 이전에 저장된 파일은 그대로 남는다.
        엑셀이 대상 파일을 열어 두었으면 os.replace가 PermissionError(공유 위반)를 내므로
        _save_with_retry의 재시도·임시저장 대체가 그대로 동작한다. 어떤 예외든 임시 파일은 지우고 다시 올린다.
        """
        target = Path(target)
        tmp = target.with_name(target.name + ".tmp")
        try:
            self._wb.save(tmp)
            os.replace(tmp, target)
        except BaseException:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                logger.warning("임시 저장 파일을 지우지 못했습니다: %s", tmp)
            raise

    def _save_with_retry(self, target: Path) -> None:
        for attempt in range(1, _SAVE_ATTEMPTS + 1):
            try:
                self._write_file(target)
                return
            except PermissionError:
                logger.warning("엑셀 저장 거부(%d/%d): %s", attempt, _SAVE_ATTEMPTS, target)
                if attempt == _SAVE_ATTEMPTS:
                    raise
                self._sleep(_RETRY_INTERVAL_SEC)

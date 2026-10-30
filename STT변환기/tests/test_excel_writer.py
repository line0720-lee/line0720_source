"""excel_writer 테스트: 저장된 xlsx를 openpyxl로 다시 열어 파일명·서식·값·초과 처리·저장 실패 대체를 확인한다."""
import ctypes
import shutil
import sys
import threading
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import excel_writer
from batch_runner import RowWriter, run_batch
from excel_writer import CELL_LIMIT, OVERFLOW_NOTE, ExcelWriter
from records import ResultRow

STARTED = datetime(2026, 10, 28, 9, 30, 15)
HEADERS = ["번호", "파일명", "변환 내용", "재생 시간", "처리 시간(초)", "상태", "변환 일시"]


def make_row(index=1, name="회의녹음_01.wav", text="오늘 회의 안건은", duration_sec=754.4,
             elapsed_sec=85.4, status="완료", finished_at=datetime(2026, 10, 28, 9, 31, 40)) -> ResultRow:
    return ResultRow(index, name, text, duration_sec, elapsed_sec, status, finished_at)


def no_sleep(_seconds: float) -> None:
    """재시도 간격을 건너뛰는 테스트용 sleep."""


def make_writer(tmp_path: Path, on_warning=None) -> ExcelWriter:
    return ExcelWriter(tmp_path, STARTED, on_warning, sleep=no_sleep)


def open_sheet(path: Path, name: str = "STT결과"):
    return load_workbook(path)[name]


def test_filename_and_location(tmp_path):
    writer = make_writer(tmp_path)
    assert writer.path == tmp_path / "STT_결과_20261028_093015.xlsx"
    assert writer.path.exists()  # 첫 파일이 끝나기 전에 중지해도 머리글만 있는 파일이 남는다


def test_satisfies_row_writer_protocol(tmp_path):
    writer: RowWriter = make_writer(tmp_path)
    assert isinstance(writer.path, Path)
    assert callable(writer.add_row)


def test_header_and_format(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row())
    wb = load_workbook(writer.path)
    assert wb.sheetnames == ["STT결과"]
    ws = wb["STT결과"]
    assert [ws.cell(1, c).value for c in range(1, 8)] == HEADERS
    for col in range(1, 8):
        cell = ws.cell(1, col)
        assert cell.font.bold
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.rgb.endswith("D9D9D9")
    assert ws.freeze_panes == "A2"
    widths = {c: ws.column_dimensions[c].width for c in "ABCDEFG"}
    assert widths == {"A": 6, "B": 40, "C": 100, "D": 12, "E": 14, "F": 28, "G": 20}
    assert ws["C2"].alignment.wrap_text is True
    for col in "ABCDEFG":
        assert ws[f"{col}2"].alignment.vertical == "top"


def test_row_values(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row())
    writer.add_row(make_row(2, "깨짐.wav", "", None, 0.2, "실패(파일 손상)", datetime(2026, 10, 28, 9, 32, 0)))
    ws = open_sheet(writer.path)
    assert [ws.cell(2, c).value for c in range(1, 8)] == [
        1, "회의녹음_01.wav", "오늘 회의 안건은", "00:12:34", 85, "완료", "2026-10-28 09:31:40",
    ]
    assert ws["E2"].value == 85 and isinstance(ws["E2"].value, int)
    assert not (ws["F2"].font.color and str(ws["F2"].font.color.rgb).endswith("FF0000"))
    # 실패 행: 재생 시간은 빈칸, 상태만 빨간 글씨
    assert ws["D3"].value is None
    assert ws["F3"].font.color.rgb.endswith("FF0000")
    assert not (ws["B3"].font.color and str(ws["B3"].font.color.rgb).endswith("FF0000"))


def test_duration_formatting(tmp_path):
    writer = make_writer(tmp_path)
    for i, sec in enumerate([0.0, 15.4, 59.6, 3899.0, 25 * 3600.0], start=1):
        writer.add_row(make_row(i, f"f{i}.wav", "x", sec))
    ws = open_sheet(writer.path)
    assert [ws.cell(r, 4).value for r in range(2, 7)] == ["00:00:00", "00:00:15", "00:01:00", "01:04:59", "25:00:00"]


def test_text_starting_with_equals_stays_text(tmp_path):
    """'='로 시작하는 변환 결과가 수식으로 해석되면 안 된다."""
    writer = make_writer(tmp_path)
    writer.add_row(make_row(text="=1+1 이라고 말했다"))
    assert open_sheet(writer.path)["C2"].value == "=1+1 이라고 말했다"
    assert open_sheet(writer.path)["C2"].data_type == "s"


def test_illegal_control_characters_do_not_break_save(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row(text="가\x00나\x0b다"))
    assert open_sheet(writer.path)["C2"].value == "가나다"


def test_saved_after_each_row(tmp_path):
    writer = make_writer(tmp_path)
    assert open_sheet(writer.path).max_row == 1
    writer.add_row(make_row())
    ws = open_sheet(writer.path)  # writer를 닫지 않은 상태에서 열린다
    assert ws.max_row == 2
    writer.add_row(make_row(2, "b.wav"))
    assert open_sheet(writer.path).max_row == 3


def test_hundred_rows(tmp_path):
    writer = make_writer(tmp_path)
    for i in range(1, 101):
        writer.add_row(make_row(i, f"파일{i:03d}.wav", f"내용 {i}"))
    ws = open_sheet(writer.path)
    assert ws.max_row == 101
    assert [ws.cell(r, 1).value for r in range(2, 102)] == list(range(1, 101))


def test_overflow_cell_sheet_txt(tmp_path):
    text = "".join(chr(0xAC00 + i % 500) for i in range(70_000))
    writer = make_writer(tmp_path)
    writer.add_row(make_row(7, "2팀\\회의녹음_01.wav", text))
    wb = load_workbook(writer.path)
    cell = wb["STT결과"]["C2"].value
    assert len(cell) <= CELL_LIMIT
    assert cell.endswith(OVERFLOW_NOTE)
    assert cell.startswith(text[:100])

    full = wb["전문"]
    assert [full.cell(1, c).value for c in range(1, 5)] == ["번호", "파일명", "조각", "내용"]
    assert full["A1"].font.bold and full["A1"].fill.fgColor.rgb.endswith("D9D9D9")
    assert full.freeze_panes == "A2"
    assert full.column_dimensions["D"].width == 100
    assert full["D2"].alignment.wrap_text is True and full["D2"].alignment.vertical == "top"
    rows = [[full.cell(r, c).value for c in range(1, 5)] for r in range(2, full.max_row + 1)]
    assert len(rows) == 3
    assert [(r[0], r[1], r[2]) for r in rows] == [(7, "2팀\\회의녹음_01.wav", n) for n in (1, 2, 3)]
    assert all(len(r[3]) <= CELL_LIMIT for r in rows)
    assert "".join(r[3] for r in rows) == text

    txt = tmp_path / "2팀_회의녹음_01.txt"
    assert txt.read_text(encoding="utf-8-sig") == text
    assert txt.read_bytes().startswith(b"\xef\xbb\xbf")


def test_no_overflow_means_no_fulltext_sheet_or_txt(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row(text="가" * CELL_LIMIT))  # 정확히 한계 길이는 초과가 아니다
    wb = load_workbook(writer.path)
    assert wb.sheetnames == ["STT결과"]
    assert len(wb["STT결과"]["C2"].value) == CELL_LIMIT
    assert list(tmp_path.glob("*.txt")) == []


def test_overflow_txt_name_uppercase_wav_suffix(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row(name="녹음.WAV", text="가" * (CELL_LIMIT + 1)))
    assert (tmp_path / "녹음.txt").exists()


def test_overflow_txt_collision_uses_timestamp_name(tmp_path):
    """result 폴더를 여러 번 쓰므로 같은 이름의 txt가 있으면 덮어쓰지 않고 실행 시각을 붙인다."""
    old = tmp_path / "회의.txt"
    old.write_text("이전 실행 결과", encoding="utf-8-sig")
    text = "나" * (CELL_LIMIT + 5)
    writer = make_writer(tmp_path)
    writer.add_row(make_row(name="회의.wav", text=text))
    assert old.read_text(encoding="utf-8-sig") == "이전 실행 결과"
    assert (tmp_path / "회의_20261028_093015.txt").read_text(encoding="utf-8-sig") == text


def test_overflow_txt_collision_within_same_run(tmp_path):
    """'2팀\\a.wav'와 '2팀_a.wav'처럼 같은 txt 이름이 한 실행 안에서 겹쳐도 서로 덮어쓰지 않는다."""
    writer = make_writer(tmp_path)
    writer.add_row(make_row(1, "2팀\\a.wav", "가" * (CELL_LIMIT + 1)))
    writer.add_row(make_row(2, "2팀_a.wav", "나" * (CELL_LIMIT + 1)))
    writer.add_row(make_row(3, "2팀\\a.wav", "다" * (CELL_LIMIT + 1)))
    contents = sorted(p.read_text(encoding="utf-8-sig")[0] for p in tmp_path.glob("*.txt"))
    assert contents == ["가", "나", "다"]


def test_second_overflow_appends_to_existing_fulltext_sheet(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row(1, "a.wav", "가" * (CELL_LIMIT + 1)))
    writer.add_row(make_row(2, "b.wav", "나" * (CELL_LIMIT * 2 + 1)))
    full = load_workbook(writer.path)["전문"]
    assert [(full.cell(r, 1).value, full.cell(r, 3).value) for r in range(2, full.max_row + 1)] == [
        (1, 1), (1, 2), (2, 1), (2, 2), (2, 3),
    ]


def test_overflow_txt_equals_fulltext_sheet_when_control_chars(tmp_path):
    """제어 문자가 섞여도 txt 내용은 '전문' 시트 조각을 이어 붙인 것과 같다(둘 다 제거본)."""
    text = "가\x01" * (CELL_LIMIT + 10) + "나" * 100  # 제어 문자를 빼도 한도를 넘는다
    writer = make_writer(tmp_path)
    writer.add_row(make_row(1, "제어.wav", text))
    full = load_workbook(writer.path)["전문"]
    joined = "".join(full.cell(r, 4).value for r in range(2, full.max_row + 1))
    assert "\x01" not in joined
    assert (tmp_path / "제어.txt").read_text(encoding="utf-8-sig") == joined


def test_overflow_txt_write_failure_still_records_row(tmp_path, caplog):
    """txt를 쓰지 못해도(경로 너무 김·디스크 부족 등) 일괄 처리를 멈추지 않고 행을 기록한다."""
    import logging

    warnings: list[str] = []
    writer = make_writer(tmp_path, warnings.append)
    writer._free_txt_path = lambda name: tmp_path / "없는 폴더" / "회의.txt"  # 쓰기에서 FileNotFoundError(OSError)
    text = "가" * (CELL_LIMIT + 10)
    with caplog.at_level(logging.ERROR, logger="excel_writer"):
        writer.add_row(make_row(3, "회의.wav", text))
    writer.add_row(make_row(4, "다음.wav", "다음 파일"))

    wb = load_workbook(writer.path)
    ws = wb["STT결과"]
    assert [ws.cell(r, 2).value for r in (2, 3)] == ["회의.wav", "다음.wav"]
    shown = ws["C2"].value
    assert len(shown) <= CELL_LIMIT
    assert shown.endswith(excel_writer.OVERFLOW_NOTE_NO_TXT)
    assert "txt" in excel_writer.OVERFLOW_NOTE_NO_TXT and "전문" in excel_writer.OVERFLOW_NOTE_NO_TXT
    full = wb["전문"]
    assert "".join(full.cell(r, 4).value for r in range(2, full.max_row + 1)) == text  # 시트에는 전체가 있다
    assert len(warnings) == 1
    assert "전문 txt 파일을 저장하지 못했습니다" in warnings[0] and "회의.txt" in warnings[0]
    assert "'전문' 시트에는 저장됨" in warnings[0]
    assert any(r.exc_info for r in caplog.records)
    assert list(tmp_path.glob("*.txt")) == []


def test_overflow_note_mentions_txt_when_txt_saved(tmp_path):
    writer = make_writer(tmp_path)
    writer.add_row(make_row(text="가" * (CELL_LIMIT + 1)))
    assert open_sheet(writer.path)["C2"].value.endswith(OVERFLOW_NOTE)


# --- 원자적 저장(임시 파일 → os.replace) ------------------------------------


def test_save_goes_through_tmp_file_then_replace(tmp_path, monkeypatch):
    writer = make_writer(tmp_path)
    saved: list[Path] = []
    real_save = writer._wb.save

    def spy(target):
        saved.append(Path(target))
        real_save(target)

    monkeypatch.setattr(writer._wb, "save", spy)
    writer.add_row(make_row())
    assert saved == [tmp_path / "STT_결과_20261028_093015.xlsx.tmp"]
    assert open_sheet(writer.path).max_row == 2
    assert list(tmp_path.glob("*.tmp")) == []


@pytest.mark.parametrize("error", [OSError(28, "No space left on device"), KeyboardInterrupt()])
def test_failed_save_keeps_previous_file_and_leaves_no_tmp(tmp_path, monkeypatch, error):
    """저장 도중 실패(디스크 부족·강제 중단)해도 이미 저장된 행은 그대로 남고 임시 파일도 남지 않는다."""
    writer = make_writer(tmp_path)
    writer.add_row(make_row(1, "a.wav"))

    def broken_save(target):
        Path(target).write_bytes(b"PK\x03\x04 truncated")  # 쓰다 만 상태
        raise error

    monkeypatch.setattr(writer._wb, "save", broken_save)
    with pytest.raises(type(error)):
        writer.add_row(make_row(2, "b.wav"))

    ws = open_sheet(writer.path)
    assert ws.max_row == 2 and ws["B2"].value == "a.wav"
    assert list(tmp_path.glob("*.tmp")) == []


def test_locked_target_on_replace_retries_then_falls_back_without_tmp(tmp_path, monkeypatch):
    """엑셀이 연 파일 위로 os.replace하면 PermissionError → 3회 시도 후 임시저장. 임시 파일은 남지 않는다."""
    warnings: list[str] = []
    sleeps: list[float] = []
    writer = ExcelWriter(tmp_path, STARTED, warnings.append, sleep=sleeps.append)
    original = writer.path
    replaced: list[Path] = []
    real_replace = excel_writer.os.replace

    def locked_replace(src, dst):
        replaced.append(Path(dst))
        if Path(dst) == original:
            raise PermissionError(13, "The process cannot access the file", str(dst))
        real_replace(src, dst)

    monkeypatch.setattr(excel_writer.os, "replace", locked_replace)
    writer.add_row(make_row())

    fallback = tmp_path / "STT_결과_20261028_093015_임시저장.xlsx"
    assert replaced == [original] * 3 + [fallback]
    assert sleeps == [1, 1]
    assert writer.path == fallback and len(warnings) == 1
    assert open_sheet(fallback).max_row == 2
    assert list(tmp_path.glob("*.tmp")) == []


def _lock_on_original(monkeypatch, writer: ExcelWriter, original: Path, calls: list):
    real_write = writer._write_file

    def fake_write(target):
        calls.append(Path(target))
        if Path(target) == original:
            raise PermissionError(13, "Permission denied", str(target))
        real_write(target)

    monkeypatch.setattr(writer, "_write_file", fake_write)


def test_save_falls_back_when_locked(tmp_path, monkeypatch):
    warnings: list[str] = []
    sleeps: list[float] = []
    writer = ExcelWriter(tmp_path, STARTED, warnings.append, sleep=sleeps.append)
    original = writer.path
    calls: list[Path] = []
    _lock_on_original(monkeypatch, writer, original, calls)

    writer.add_row(make_row())

    fallback = tmp_path / "STT_결과_20261028_093015_임시저장.xlsx"
    assert writer.path == fallback
    assert calls == [original] * 3 + [fallback]  # 3회 시도 후 대체 경로 저장
    assert sleeps == [1, 1]  # 시도 사이에만 1초 간격
    assert len(warnings) == 1
    assert "STT_결과_20261028_093015_임시저장.xlsx" in warnings[0]
    assert "열려 있어" in warnings[0]
    assert open_sheet(fallback).max_row == 2

    # 이후 행은 알림 없이 새 파일에 이어서 저장된다
    calls.clear()
    writer.add_row(make_row(2, "b.wav"))
    assert calls == [fallback]
    assert len(warnings) == 1
    assert open_sheet(fallback).max_row == 3


def test_recovers_without_fallback_when_lock_clears_within_retries(tmp_path, monkeypatch):
    warnings: list[str] = []
    writer = ExcelWriter(tmp_path, STARTED, warnings.append, sleep=no_sleep)
    original = writer.path
    real_write = writer._write_file
    attempts = {"n": 0}

    def flaky_write(target):
        attempts["n"] += 1
        if attempts["n"] < 3:
            raise PermissionError(13, "Permission denied")
        real_write(target)

    monkeypatch.setattr(writer, "_write_file", flaky_write)
    writer.add_row(make_row())
    assert writer.path == original and warnings == []
    assert open_sheet(original).max_row == 2


def test_fallback_save_failure_propagates(tmp_path, monkeypatch):
    writer = make_writer(tmp_path)

    def always_denied(_target):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(writer, "_write_file", always_denied)
    with pytest.raises(PermissionError):
        writer.add_row(make_row())


def test_non_permission_oserror_propagates_immediately(tmp_path, monkeypatch):
    sleeps: list[float] = []
    writer = ExcelWriter(tmp_path, STARTED, None, sleep=sleeps.append)
    calls: list[Path] = []

    def disk_full(target):
        calls.append(Path(target))
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(writer, "_write_file", disk_full)
    with pytest.raises(OSError):
        writer.add_row(make_row())
    assert len(calls) == 1 and sleeps == [] and writer.path.name == "STT_결과_20261028_093015.xlsx"


def test_fallback_logs_warning(tmp_path, monkeypatch, caplog):
    import logging

    writer = make_writer(tmp_path)
    _lock_on_original(monkeypatch, writer, writer.path, [])
    with caplog.at_level(logging.WARNING, logger="excel_writer"):
        writer.add_row(make_row())
    assert any("임시저장" in r.getMessage() for r in caplog.records)


def test_works_without_on_warning_callback(tmp_path, monkeypatch):
    writer = make_writer(tmp_path, on_warning=None)
    _lock_on_original(monkeypatch, writer, writer.path, [])
    writer.add_row(make_row())
    assert writer.path.name.endswith("_임시저장.xlsx")


@pytest.mark.skipif(sys.platform != "win32", reason="Windows 공유 잠금 전용")
def test_save_falls_back_with_real_windows_lock(tmp_path):
    """엑셀이 파일을 열어 둔 상태를 CreateFileW(공유 모드 0)로 재현한다."""
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = [
        ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
        ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
    ]
    kernel32.CreateFileW.restype = ctypes.c_void_p
    kernel32.CloseHandle.argtypes = [ctypes.c_void_p]

    warnings: list[str] = []
    writer = ExcelWriter(tmp_path, STARTED, warnings.append, sleep=no_sleep)
    original = writer.path
    GENERIC_READ, OPEN_EXISTING = 0x80000000, 3
    handle = kernel32.CreateFileW(str(original), GENERIC_READ, 0, None, OPEN_EXISTING, 0x80, None)
    assert handle not in (None, ctypes.c_void_p(-1).value), ctypes.get_last_error()
    try:
        writer.add_row(make_row())
    finally:
        kernel32.CloseHandle(handle)

    fallback = tmp_path / "STT_결과_20261028_093015_임시저장.xlsx"
    assert writer.path == fallback and fallback.exists()
    assert len(warnings) == 1
    assert open_sheet(fallback).max_row == 2
    assert open_sheet(original).max_row == 1  # 잠겨 있던 원본은 건드리지 못했다
    assert list(tmp_path.glob("*.tmp")) == []  # os.replace가 거부된 임시 파일도 지웠다


@pytest.mark.model
def test_real_batch_writes_xlsx(samples_dir, small_engine, tmp_path):
    work = tmp_path / "samples"
    shutil.copytree(samples_dir, work)
    holder: dict[str, ExcelWriter] = {}

    def factory(result_dir, started_at):
        holder["writer"] = ExcelWriter(result_dir, started_at)
        return holder["writer"]

    summary = run_batch(work, True, small_engine, factory, lambda p: None, threading.Event())
    assert summary.excel_path == holder["writer"].path
    assert summary.excel_path.parent == work / "result"
    assert summary.excel_path.name.startswith("STT_결과_") and summary.excel_path.suffix == ".xlsx"

    ws = open_sheet(summary.excel_path)
    def is_red(cell) -> bool:
        return bool(cell.font.color and str(cell.font.color.rgb).endswith("FF0000"))

    for r in range(1, ws.max_row + 1):
        print([ws.cell(r, c).value for c in range(1, 8)], "| F빨강:", is_red(ws.cell(r, 6)))

    assert ws.max_row == 8  # 머리글 + 7건
    assert [ws.cell(1, c).value for c in range(1, 8)] == HEADERS
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("D9D9D9") and ws.freeze_panes == "A2"
    rows = {ws.cell(r, 2).value: r for r in range(2, ws.max_row + 1)}
    status = lambda name: ws.cell(rows[name], 6).value  # noqa: E731
    for ok in ("기본_16k.wav", "전화_8k.wav", "스테레오_44k.wav", "하위 폴더 (2팀)\\회의 #1.wav"):
        assert status(ok) == "완료"
    assert status("무음.wav") == "완료(인식된 음성 없음)"
    assert status("빈파일.wav") == "실패(파일 손상)"
    assert status("가짜.wav") == "실패(파일 손상)"

    red_rows = {ws.cell(r, 2).value for r in range(2, ws.max_row + 1) if is_red(ws.cell(r, 6))}
    assert red_rows == {"빈파일.wav", "가짜.wav"}
    assert [ws.cell(r, 1).value for r in range(2, 9)] == list(range(1, 8))
    for name in ("기본_16k.wav", "전화_8k.wav", "스테레오_44k.wav"):
        assert ws.cell(rows[name], 4).value == "00:00:15", name
    assert ws.cell(rows["빈파일.wav"], 4).value is None

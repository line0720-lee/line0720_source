"""main 명령줄(--no-gui) 모드 테스트: 종료 코드와 결과 엑셀을 확인한다."""
import logging
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import config
import main
from batch_runner import OutputNotWritableError
from config import AppConfig, ConfigError
from stt_engine import TranscribeResult


@pytest.fixture(autouse=True)
def tmp_logging(tmp_path, monkeypatch):
    """main()이 프로젝트 logs 폴더 대신 임시 폴더에 로그를 남기게 하고, 테스트 뒤 핸들러를 정리한다."""
    root = logging.getLogger()
    before = list(root.handlers)
    level = root.level
    log_base = tmp_path / "앱"
    monkeypatch.setattr(main, "setup_logging", lambda: config.setup_logging(log_base))
    yield log_base / "logs"
    for handler in list(root.handlers):
        if handler not in before:
            root.removeHandler(handler)
            handler.close()
    root.setLevel(level)


def _make_folder(tmp_path: Path, *names: str) -> Path:
    folder = tmp_path / "녹음 폴더 (테스트)"
    folder.mkdir()
    for name in names:
        (folder / name).write_bytes(b"RIFF")
    return folder


def _fake_config(tmp_path: Path) -> AppConfig:
    """model.bin만 있는 가짜 모델 폴더를 가리키는 설정."""
    model_dir = tmp_path / "models" / "small"
    model_dir.mkdir(parents=True)
    (model_dir / "model.bin").write_bytes(b"x")
    return AppConfig("small", model_dir, "ko", 1, tmp_path / "glossary.txt")


class FakeEngine:
    def __init__(self, cfg):
        self.cfg = cfg

    def transcribe(self, wav, on_progress=None):
        if on_progress is not None:
            on_progress(1.0)
        return TranscribeResult(text=f"{Path(wav).name} 내용", duration_sec=3.0, elapsed_sec=0.1)


def _data_rows(folder: Path) -> list[tuple]:
    files = list((folder / "result").glob("*.xlsx"))
    assert len(files) == 1
    wb = load_workbook(files[0])
    try:
        return list(wb["STT결과"].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


# --------------------------------------------------------------- 실제 small 모델


@pytest.mark.model
def test_cli_converts_samples(samples_dir, small_cfg, tmp_path, tmp_logging):
    folder = tmp_path / "녹음 폴더 (테스트)"
    folder.mkdir()
    for name in ("기본_16k.wav", "빈파일.wav"):
        shutil.copy(samples_dir / name, folder / name)

    code = main.main(["--folder", str(folder), "--no-gui", "--model", "small"])

    assert code == 1  # 빈파일.wav가 실패하므로 1
    rows = _data_rows(folder)
    assert len(rows) == 2
    status = {row[1]: row[5] for row in rows}
    assert status["기본_16k.wav"] == "완료"
    assert status["빈파일.wav"] == "실패(파일 손상)"
    log_text = "".join(p.read_text(encoding="utf-8") for p in tmp_logging.glob("*.log"))
    assert "일괄 처리 시작" in log_text


@pytest.mark.model
def test_cli_no_wavs_returns_2(tmp_path):
    folder = _make_folder(tmp_path)
    assert main.main(["--folder", str(folder), "--no-gui", "--model", "small"]) == 2
    assert not (folder / "result").exists()


# --------------------------------------------------------------- 모델 없이 확인하는 경우


def test_cli_all_success_returns_0(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "회의 1.wav", "회의 2.wav")
    monkeypatch.setattr(main, "load_config", lambda: _fake_config(tmp_path))
    monkeypatch.setattr(main, "SttEngine", FakeEngine)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 0
    rows = _data_rows(folder)
    assert [row[5] for row in rows] == ["완료", "완료"]


def test_cli_subfolders_option(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "위.wav")
    (folder / "하위 (2팀)").mkdir()
    (folder / "하위 (2팀)" / "아래.wav").write_bytes(b"RIFF")
    monkeypatch.setattr(main, "load_config", lambda: _fake_config(tmp_path))
    monkeypatch.setattr(main, "SttEngine", FakeEngine)
    assert main.main(["--folder", str(folder), "--no-gui", "--subfolders"]) == 0
    assert [row[1] for row in _data_rows(folder)] == ["위.wav", "하위 (2팀)\\아래.wav"]


def test_cli_requires_folder():
    assert main.main(["--no-gui"]) == 2


def test_cli_missing_folder_returns_2(tmp_path):
    assert main.main(["--folder", str(tmp_path / "없는 폴더"), "--no-gui"]) == 2


def test_cli_config_error_returns_2(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "a.wav")

    def broken():
        raise ConfigError("설정 파일(config.ini)을 찾을 수 없습니다")

    monkeypatch.setattr(main, "load_config", broken)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2
    assert not (folder / "result").exists()


def test_cli_model_missing_returns_2(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "a.wav")
    cfg = AppConfig("medium", tmp_path / "models" / "medium", "ko", 1, tmp_path / "glossary.txt")
    monkeypatch.setattr(main, "load_config", lambda: cfg)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2
    assert not (folder / "result").exists()


def test_cli_model_load_failure_returns_2(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "a.wav")
    monkeypatch.setattr(main, "load_config", lambda: _fake_config(tmp_path))

    def failing_engine(cfg):
        raise RuntimeError("모델을 불러오지 못했습니다: 테스트")

    monkeypatch.setattr(main, "SttEngine", failing_engine)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2


def test_cli_output_not_writable_returns_2(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "a.wav")
    monkeypatch.setattr(main, "load_config", lambda: _fake_config(tmp_path))
    monkeypatch.setattr(main, "SttEngine", FakeEngine)

    def not_writable(*args, **kwargs):
        raise OutputNotWritableError("결과 폴더에 쓸 수 없습니다")

    monkeypatch.setattr(main, "run_batch", not_writable)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2


def _use_fakes(tmp_path, monkeypatch, engine=FakeEngine):
    monkeypatch.setattr(main, "load_config", lambda: _fake_config(tmp_path))
    monkeypatch.setattr(main, "SttEngine", engine)


# --------------------------------------------------------------- 위치 인자(폴더를 exe에 끌어다 놓기)


def test_cli_positional_folder(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "회의 1.wav")
    _use_fakes(tmp_path, monkeypatch)
    assert main.main([str(folder), "--no-gui"]) == 0
    assert [row[1] for row in _data_rows(folder)] == ["회의 1.wav"]


def test_cli_folder_option_wins_over_positional(tmp_path, monkeypatch):
    folder = _make_folder(tmp_path, "회의 1.wav")
    _use_fakes(tmp_path, monkeypatch)
    assert main.main([str(tmp_path / "없는 폴더"), "--folder", str(folder), "--no-gui"]) == 0
    assert len(_data_rows(folder)) == 1


def test_cli_positional_missing_folder_returns_2(tmp_path):
    assert main.main([str(tmp_path / "없는 폴더"), "--no-gui"]) == 2


@pytest.mark.parametrize("argv", [["--bogus"], ["폴더1", "폴더2"], ["--model", "large"]])
def test_gui_mode_usage_error_shows_window(argv, monkeypatch):
    shown = []
    monkeypatch.setattr(main, "_show_startup_error", shown.append)
    assert main.main(argv) == 2
    assert len(shown) == 1 and "명령줄 인자가 잘못되었습니다" in shown[0]


def test_cli_usage_error_goes_to_stderr(monkeypatch, capsys):
    shown = []
    monkeypatch.setattr(main, "_show_startup_error", shown.append)
    assert main.main(["--no-gui", "--bogus"]) == 2
    assert shown == []
    assert "오류:" in capsys.readouterr().err


# --------------------------------------------------------------- Ctrl+C, 일괄 처리 중 오류


def test_cli_ctrl_c_during_model_load_returns_2(tmp_path, monkeypatch, capsys):
    folder = _make_folder(tmp_path, "a.wav")

    def interrupted(cfg):
        raise KeyboardInterrupt

    _use_fakes(tmp_path, monkeypatch, engine=interrupted)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2
    assert "사용자가 중단했습니다" in capsys.readouterr().out


def test_cli_ctrl_c_during_batch_returns_1(tmp_path, monkeypatch, capsys):
    folder = _make_folder(tmp_path, "1.wav", "2.wav", "3.wav")

    class InterruptOnSecond(FakeEngine):
        def transcribe(self, wav, on_progress=None):
            if Path(wav).name == "2.wav":
                raise KeyboardInterrupt
            return super().transcribe(wav, on_progress)

    _use_fakes(tmp_path, monkeypatch, engine=InterruptOnSecond)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 1
    out = capsys.readouterr().out
    assert "사용자가 중단했습니다" in out and "STT_결과_" in out
    assert [row[1] for row in _data_rows(folder)] == ["1.wav"]


def test_cli_memory_error_outside_file_gives_small_advice(tmp_path, monkeypatch, capsys):
    folder = _make_folder(tmp_path, "a.wav")
    _use_fakes(tmp_path, monkeypatch)

    def out_of_memory(*args, **kwargs):
        raise MemoryError

    monkeypatch.setattr(main, "run_batch", out_of_memory)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2  # 저장된 행 없음
    assert "config.ini의 name을 small로" in capsys.readouterr().out


def test_cli_save_failure_names_partial_file(tmp_path, monkeypatch, capsys):
    folder = _make_folder(tmp_path, "1.wav", "2.wav")
    _use_fakes(tmp_path, monkeypatch)

    class FailingSecondSave(main.ExcelWriter):
        def add_row(self, row):
            if row.index == 2:
                raise PermissionError("저장 거부(테스트)")
            super().add_row(row)

    monkeypatch.setattr(main, "ExcelWriter", FailingSecondSave)
    assert main.main(["--folder", str(folder), "--no-gui"]) == 1  # 1건은 저장됨
    out = capsys.readouterr().out
    xlsx = next((folder / "result").glob("*.xlsx"))
    assert "변환 중 오류가 발생했습니다" in out and xlsx.name in out


class _FailingSaveFrom(main.ExcelWriter):
    """index번째 행부터 저장이 거부되는 writer(엑셀 저장 실패로 일괄 처리가 멈추는 경우)."""

    fail_from = 1

    def add_row(self, row):
        if row.index >= self.fail_from:
            raise PermissionError("저장 거부(테스트)")
        super().add_row(row)


def test_cli_save_failure_before_any_row_does_not_claim_saved_results(tmp_path, monkeypatch, capsys):
    folder = _make_folder(tmp_path, "1.wav", "2.wav")
    _use_fakes(tmp_path, monkeypatch)
    monkeypatch.setattr(main, "ExcelWriter", type("FailFirst", (_FailingSaveFrom,), {"fail_from": 1}))
    assert main.main(["--folder", str(folder), "--no-gui"]) == 2  # 저장된 행 없음
    out = capsys.readouterr().out
    assert "변환 중 오류가 발생했습니다" in out
    assert "지금까지의 결과" not in out


# --------------------------------------------------------------- 화면 모드 작업 스레드(App._work, 창 없이)


def _work_events(tmp_path, monkeypatch, writer_cls, *names) -> list[tuple]:
    """Tk 창을 만들지 않고 App._work만 실행해 queue에 넣은 이벤트를 돌려준다."""
    import queue
    import threading

    folder = _make_folder(tmp_path, *names)
    monkeypatch.setattr(main, "ExcelWriter", writer_cls)
    app = object.__new__(main.App)  # __init__(창 생성)은 건너뛴다. _work가 쓰는 속성만 채운다
    app.events = queue.Queue()
    app.engine = FakeEngine(None)
    app.cfg = _fake_config(tmp_path)
    app.log_dir = None
    app._work(folder, False, threading.Event())
    events = []
    while not app.events.empty():
        events.append(app.events.get_nowait())
    return events


def test_gui_work_save_failure_after_rows_points_to_partial_file(tmp_path, monkeypatch):
    events = _work_events(tmp_path, monkeypatch, type("FailSecond", (_FailingSaveFrom,), {"fail_from": 2}),
                          "1.wav", "2.wav")
    kinds = [kind for kind, _ in events]
    assert kinds.count("row") == 1 and kinds[-1] == "finished"
    paths = [payload for kind, payload in events if kind == "result_path"]
    assert len(paths) == 1 and paths[0].name.startswith("STT_결과_")
    error = next(payload for kind, payload in events if kind == "error")
    assert "지금까지의 결과는 저장되어 있습니다" in error and paths[0].name in error


def test_gui_work_save_failure_before_any_row_has_no_partial_claim(tmp_path, monkeypatch):
    events = _work_events(tmp_path, monkeypatch, type("FailFirst", (_FailingSaveFrom,), {"fail_from": 1}),
                          "1.wav", "2.wav")
    kinds = [kind for kind, _ in events]
    assert "row" not in kinds and "result_path" not in kinds
    error = next(payload for kind, payload in events if kind == "error")
    assert "변환 중 오류가 발생했습니다" in error and "지금까지의 결과" not in error


def test_gui_work_success_puts_rows_and_summary(tmp_path, monkeypatch):
    events = _work_events(tmp_path, monkeypatch, main.ExcelWriter, "1.wav", "2.wav")
    kinds = [kind for kind, _ in events]
    assert kinds.count("row") == 2 and "summary" in kinds and "error" not in kinds


def test_model_override_changes_model_path(tmp_path, monkeypatch):
    base = _fake_config(tmp_path)
    medium = AppConfig("medium", base.model_path.parent / "medium", "ko", 1, base.glossary_path)
    monkeypatch.setattr(main, "load_config", lambda: medium)
    cfg = main.load_app_config("small")
    assert cfg.model_name == "small"
    assert cfg.model_path == tmp_path / "models" / "small"

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import httpx
import json
import io
import re
import os
import uuid

import fitz  # PyMuPDF

# ── 자동 로딩 대상 폴더 ──────────────────────────────────
DATA_DIR   = r"C:\Claude(바이브코딩)\Study\IT_Error\Data"
IMAGES_DIR = os.path.join(DATA_DIR, "images")
os.makedirs(IMAGES_DIR, exist_ok=True)

ALLOWED_IMG_EXTS = {"png", "jpg", "jpeg", "gif", "webp"}

TEXT_EXTS_BE = {
    'txt','log','csv','tsv','json','jsonl','md','markdown',
    'xml','yaml','yml','ini','cfg','conf',
    'py','js','ts','cs','java','sql','sh','bat','ps1',
}

SECTION_ORDER = [
    "Ⅰ. PC·단말기 관련",
    "Ⅱ. 네트워크 관련",
    "Ⅲ. 프린터·주변기기 관련",
    "Ⅳ. 병원정보시스템(OCS·EMR·PACS 등) 관련",
    "Ⅴ. 서버·전원 관련",
    "Ⅵ. 보안 관련",
]

def _find_docx() -> str | None:
    """DATA_DIR 에서 첫 번째 .docx 파일 경로 반환."""
    for fname in sorted(os.listdir(DATA_DIR)):
        if fname.endswith('.docx') and not fname.startswith(('~', '.')):
            return os.path.join(DATA_DIR, fname)
    return None

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/images", StaticFiles(directory=IMAGES_DIR), name="images")

# ── 공유 게이트웨이 주소 (모든 모델 호출은 이 변수만 사용) ──────────────
GATEWAY      = "http://121.138.151.6:11500"
OLLAMA_URL   = f"{GATEWAY}/api/chat"
OLLAMA_EMBED = f"{GATEWAY}/api/embed"
MODEL        = "qwen3.6:35b-a3b"
EMBED_MODEL  = "bge-m3"

SYSTEM_PROMPT = (
    "당신은 정보지원팀 전산장애·업무 안내 챗봇입니다. "
    "아래 참고 자료는 사용자 파일에서 질문과 관련성이 높은 구간만 선별(RAG)한 청크입니다. "
    "Q&A 항목은 '[구간 N] 파일: 파일명 · Q번호 (섹션)' 형식으로, "
    "일반 파일은 '[구간 N] 파일: 파일명 · 줄범위' 형식으로 표시됩니다.\n"
    "규칙 1 (근거 기반): 반드시 제공된 참고 자료에만 근거해 한국어로 답하고, "
    "자료에 없는 내용은 지어내지 말고 '확인되지 않습니다'라고 답한다.\n"
    "규칙 2 (보안 거절): 서버 비밀번호·관리자 계정·상세 IP 대역 같은 보안정보나 "
    "환자 개인정보를 묻는 질문에는 모델 추론 없이 '보안상 안내하지 않습니다'라고만 답한다.\n"
    "규칙 3 (출처 표시): 답변 마지막에 근거 출처를 명시한다. "
    "Q&A 파일은 '[Q번호]' 형식으로, 일반 파일은 '[파일명 시작줄~끝줄]' 형식으로 표시한다.\n"
    "규칙 4 (형식): 대응 절차가 여러 단계인 경우 반드시 각 단계를 줄바꿈으로 구분하여 작성한다. "
    "예시:\n"
    "1. 전원 케이블 연결 상태를 확인합니다.\n"
    "2. 재부팅 후 증상이 지속되면 담당자에게 연락합니다. [Q1]"
)


class ChatRequest(BaseModel):
    question: str
    context: str

class EmbedRequest(BaseModel):
    texts: list[str]

class QnaItem(BaseModel):
    section: str
    question: str
    answer_lines: list[str]
    images: list[str] = []

class SaveQnaRequest(BaseModel):
    items: list[QnaItem]


async def stream_ollama(question: str, context: str):
    payload = {
        "model": MODEL,
        "stream": True,
        "think": False,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f"/no_think\n참고 자료:\n{context}\n\n질문: {question}"},
        ],
    }
    async with httpx.AsyncClient(timeout=120) as client:
        async with client.stream("POST", OLLAMA_URL, json=payload) as resp:
            async for line in resp.aiter_lines():
                if not line:
                    continue
                try:
                    data = json.loads(line)
                    token = data.get("message", {}).get("content", "")
                    if token:
                        yield f"data: {json.dumps({'token': token})}\n\n"
                    if data.get("done"):
                        yield "data: [DONE]\n\n"
                except json.JSONDecodeError:
                    continue


@app.post("/api/chat")
async def chat(req: ChatRequest):
    return StreamingResponse(
        stream_ollama(req.question, req.context),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/embed")
async def embed(req: EmbedRequest):
    """텍스트 배열을 bge-m3 벡터로 변환."""
    async with httpx.AsyncClient(timeout=120) as client:
        res = await client.post(
            OLLAMA_EMBED,
            json={"model": EMBED_MODEL, "input": req.texts},
        )
        data = res.json()
        return {"embeddings": data.get("embeddings", [])}


@app.post("/api/parse-pdf")
async def parse_pdf(file: UploadFile = File(...)):
    """PDF → 페이지별 줄 목록 반환.
    텍스트 레이어(PyMuPDF) 우선, 빈 페이지는 OCR(Tesseract) 폴백."""
    data = await file.read()
    pages = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        for page_num, page in enumerate(doc, start=1):
            # ① 텍스트 레이어 추출
            text = page.get_text("text").strip()
            lines = [l.strip() for l in text.splitlines() if l.strip()]

            # ② 텍스트가 거의 없으면 OCR 시도
            if len(lines) < 3:
                try:
                    import pytesseract
                    from PIL import Image
                    mat = fitz.Matrix(2, 2)          # 2× 해상도
                    pix = page.get_pixmap(matrix=mat)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_text = pytesseract.image_to_string(img, lang="kor+eng")
                    ocr_lines = [l.strip() for l in ocr_text.splitlines() if l.strip()]
                    if ocr_lines:
                        lines = ocr_lines
                except Exception:
                    pass   # Tesseract 미설치 시 무시

            pages.append({"page": page_num, "lines": lines})
        doc.close()
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})
    return {"filename": file.filename, "pages": pages}


def _parse_qna_paragraphs(paragraphs) -> list[dict]:
    """docx 단락 리스트 → Q&A 항목 리스트 (공용 헬퍼)."""
    items, section, q_num, q_text, a_lines = [], "", None, "", []

    def flush():
        if q_num and a_lines:
            images, clean_lines = [], []
            for l in a_lines:
                m = re.match(r'^\[IMG:(.+)\]$', l)
                if m:
                    images.append(m.group(1))
                else:
                    clean_lines.append(l)
            text = f"Q{q_num}. {q_text}\n" + "\n".join(f"- {l}" for l in clean_lines)
            if images:
                text += f"\n(첨부 이미지 {len(images)}장)"
            items.append({
                "id":           f"Q{q_num}",
                "section":      section,
                "question":     q_text,
                "answer_lines": clean_lines,
                "images":       images,
                "text":         text,
            })

    for para in paragraphs:
        t = para.text.strip()
        if not t:
            continue
        if re.match(r'^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]', t):
            flush(); q_num = None; a_lines = []
            section = re.sub(r'\s+', ' ', t)
        elif m := re.match(r'^Q(\d+)\.\s+(.+)', t):
            flush(); q_num, q_text, a_lines = m.group(1), m.group(2).strip(), []
        elif q_num:
            if t.startswith("A."):
                line = re.sub(r'^A\.\s*-?\s*', '', t).strip()
                if line: a_lines.append(line)
            elif t.startswith("-"):
                a_lines.append(t.lstrip("- ").strip())

    flush()
    return items


@app.post("/api/parse-docx")
async def parse_docx_qna(file: UploadFile = File(...)):
    """Word Q&A 파일 업로드 → Q&A 항목 리스트 반환."""
    data = await file.read()
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(data))
        items = _parse_qna_paragraphs(doc.paragraphs)
        return {"filename": file.filename, "items": items}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})


@app.get("/api/load-data")
async def load_data():
    """DATA_DIR 폴더를 자동 파싱하여 파일별 데이터 반환."""
    if not os.path.isdir(DATA_DIR):
        return JSONResponse(status_code=404, content={"error": f"폴더 없음: {DATA_DIR}"})

    result = []
    for fname in sorted(os.listdir(DATA_DIR)):
        if fname.startswith(('~', '.')):   # Word 임시파일·숨김파일 제외
            continue
        fpath = os.path.join(DATA_DIR, fname)
        if not os.path.isfile(fpath):
            continue
        ext = fname.rsplit('.', 1)[-1].lower()

        try:
            if ext == 'docx':
                from docx import Document as DocxDocument
                doc = DocxDocument(fpath)
                items = _parse_qna_paragraphs(doc.paragraphs)
                if items:
                    result.append({"filename": fname, "type": "qna", "items": items})

            elif ext == 'pdf':
                doc = fitz.open(fpath)
                pages = []
                for page_num, page in enumerate(doc, start=1):
                    text = page.get_text("text").strip()
                    lines = [l.strip() for l in text.splitlines() if l.strip()]
                    pages.append({"page": page_num, "lines": lines})
                doc.close()
                result.append({"filename": fname, "type": "pdf", "pages": pages})

            elif ext in TEXT_EXTS_BE:
                with open(fpath, encoding='utf-8', errors='ignore') as f:
                    lines = f.read().splitlines()
                result.append({"filename": fname, "type": "text", "lines": lines})

        except Exception as e:
            result.append({"filename": fname, "type": "error", "error": str(e)})

    return {"folder": DATA_DIR, "files": result}


@app.get("/api/qna")
async def get_qna():
    """문서 편집기용: docx Q&A 항목 전체 반환 (answer_lines 포함)."""
    path = _find_docx()
    if not path:
        return JSONResponse(status_code=404, content={"error": "docx 파일 없음"})
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(path)
        items = _parse_qna_paragraphs(doc.paragraphs)
        return {"items": items, "sections": SECTION_ORDER, "filename": os.path.basename(path)}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})


@app.post("/api/qna/save")
async def save_qna(req: SaveQnaRequest):
    """편집된 Q&A 목록을 docx 파일로 재저장."""
    path = _find_docx()
    if not path:
        return JSONResponse(status_code=404, content={"error": "docx 파일 없음"})
    try:
        from docx import Document as DocxDocument

        # 섹션 정렬
        order_idx = {s: i for i, s in enumerate(SECTION_ORDER)}
        item_list = list(req.items)
        item_list.sort(key=lambda x: order_idx.get(x.section, 99))

        doc = DocxDocument()
        # 제목 문단
        doc.add_paragraph("병원 전산장비 장애 대응 Q&A")
        doc.add_paragraph("정보지원팀 실무 대응 매뉴얼 (전산장비 A/S 자체처리 기준)")

        q_num = 1
        cur_sec = None
        for item in item_list:
            if item.section != cur_sec:
                cur_sec = item.section
                doc.add_paragraph(cur_sec)
            doc.add_paragraph(f"Q{q_num}. {item.question}")
            full_lines = list(item.answer_lines) + [f"[IMG:{img}]" for img in item.images]
            if full_lines:
                doc.add_paragraph(f"A.  - {full_lines[0]}")
                for line in full_lines[1:]:
                    doc.add_paragraph(f"     - {line}")
            q_num += 1

        doc.save(path)
        return {"ok": True, "saved": os.path.basename(path), "count": len(item_list)}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})


@app.post("/api/upload-image")
async def upload_image(file: UploadFile = File(...)):
    """대응방법 첨부 이미지 업로드 → Data/images 폴더에 저장."""
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in ALLOWED_IMG_EXTS:
        return JSONResponse(status_code=400, content={"error": "지원하지 않는 이미지 형식입니다. (png/jpg/jpeg/gif/webp)"})
    safe_name = f"{uuid.uuid4().hex}.{ext}"
    data = await file.read()
    with open(os.path.join(IMAGES_DIR, safe_name), "wb") as f:
        f.write(data)
    return {"filename": safe_name, "url": f"/images/{safe_name}"}


@app.get("/health")
async def health():
    return {"status": "ok"}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
